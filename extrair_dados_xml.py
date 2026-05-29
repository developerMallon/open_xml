import os
import glob
import time
from lxml import etree
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def extrair_dados_nfe():
    start_time = time.time()
    pasta_base = 'xml'
    dados_extraidos = []
    
    # Namespaces padrão da NF-e
    ns = {'nfe': 'http://www.portalfiscal.inf.br/nfe'}
    
    if not os.path.exists(pasta_base):
        print(f"Erro: A pasta '{pasta_base}' não foi encontrada no diretório atual.")
        return

    print("Iniciando varredura recursiva de arquivos XML...")
    
    # Localizar todos os arquivos .xml recursivamente
    arquivos_xml = []
    for root_dir, _, files in os.walk(pasta_base):
        for file in files:
            if file.lower().endswith('.xml'):
                arquivos_xml.append(os.path.join(root_dir, file))
                
    total_arquivos = len(arquivos_xml)
    print(f"Encontrados {total_arquivos} arquivos XML para processamento.\n")

    if total_arquivos == 0:
        print("Nenhum arquivo XML encontrado para processar.")
        return

    arquivos_com_sucesso = 0
    arquivos_com_erro = 0

    for i, caminho_completo in enumerate(arquivos_xml, 1):
        nome_arquivo = os.path.basename(caminho_completo)
        print(f"[{i}/{total_arquivos}] Processando: {nome_arquivo}...", end="", flush=True)
        
                
        try:
            # Parse do XML usando lxml
            parser = etree.XMLParser(remove_blank_text=True)
            tree = etree.parse(caminho_completo, parser=parser)
            root = tree.getroot()
            
            # Helper para simplificar XPath e evitar exceções
            def query_xpath(xpath_expr, fallback_expr=None, is_attr=False):
                res = root.xpath(xpath_expr, namespaces=ns)
                if not res and fallback_expr:
                    res = root.xpath(fallback_expr, namespaces=ns)
                
                if res:
                    val = res[0]
                    # Se for atributo ou nó de texto
                    if hasattr(val, 'strip'):
                        return val.strip()
                    elif isinstance(val, etree._Element) and val.text:
                        return val.text.strip()
                    return str(val)
                return ''

            # Extrair chassi para garantir que é uma nota referente a VEÍCULO
            chassi_veiculo = query_xpath('//nfe:prod/nfe:veicProd/nfe:chassi/text()')

            if not chassi_veiculo:
                # Pula este arquivo
                print(" PULO (NFe sem chassi - não é veículo)")
                arquivos_com_sucesso += 1
                continue

            # Extração de campos cabeçalho / gerais da nota
            nNF = query_xpath('//nfe:ide/nfe:nNF/text()')
            tpNF = query_xpath('//nfe:ide/nfe:tpNF/text()')
            vFrete = query_xpath('//nfe:total/nfe:ICMSTot/nfe:vFrete/text()')
            modFrete = query_xpath('//nfe:transp/nfe:modFrete/text()')
            
            # volumes (pode haver mais de um, pegamos o primeiro por padrão)
            qVol = query_xpath('//nfe:transp/nfe:vol/nfe:qVol/text()')
            esp = query_xpath('//nfe:transp/nfe:vol/nfe:esp/text()')
            
            # data de emissão (com fallback para dEmi)
            dhEmi = query_xpath('//nfe:ide/nfe:dhEmi/text()', '//nfe:ide/nfe:dEmi/text()')
            
            # CNPJ / CPF emitente
            emit_cnpj = query_xpath('//nfe:emit/nfe:CNPJ/text()', '//nfe:emit/nfe:CPF/text()')
            
            # CNPJ / CPF destinatário
            dest_cnpj = query_xpath('//nfe:dest/nfe:CNPJ/text()', '//nfe:dest/nfe:CPF/text()')
            
            # natureza da operação
            natOp = query_xpath('//nfe:ide/nfe:natOp/text()')
            
            # chave de acesso da NFe
            chNFe = query_xpath('//nfe:protNFe/nfe:infProt/nfe:chNFe/text()')
            if not chNFe:
                infNFe_id = query_xpath('//nfe:infNFe/@Id')
                if infNFe_id:
                    # remove o prefixo 'NFe' se houver
                    chNFe = infNFe_id[3:] if infNFe_id.startswith('NFe') else infNFe_id

            # Tratamento de valores numéricos
            try:
                vFrete_float = float(vFrete) if vFrete else 0.0
            except ValueError:
                vFrete_float = 0.0

            try:
                qVol_float = float(qVol) if qVol else 0.0
                # Se for número inteiro puro, converte para int
                if qVol_float.is_integer():
                    qVol_float = int(qVol_float)
            except ValueError:
                qVol_float = qVol if qVol else ''

            # Extrair itens para ncm e cfop
            itens = root.xpath('//nfe:det', namespaces=ns)
            
            def adicionar_registro(ncm_val, cfop_val):
                dados_extraidos.append({
                    'Chave da NFe': chNFe,
                    'Tipo NF': tpNF,
                    'Número': nNF,
                    'Data Emissão': dhEmi,
                    'CNPJ Emitente': emit_cnpj,
                    'CNPJ Destinatário': dest_cnpj,
                    'Natureza Operação': natOp,
                    'Valor Frete': vFrete_float,
                    'Mod Frete': modFrete,
                    'Qtd Volume': qVol_float,
                    'Espécie': esp,
                    'NCM': ncm_val,
                    'CFOP': cfop_val,
                    'Arquivo Origem': nome_arquivo
                })

            if not itens:
                # Nota sem itens cadastrados
                adicionar_registro('', '')
            else:
                for item in itens:
                    # Busca de forma relativa dentro do item
                    ncm_item = item.xpath('.//nfe:prod/nfe:NCM/text()', namespaces=ns)
                    ncm_item = ncm_item[0].strip() if ncm_item else ''
                    
                    cfop_item = item.xpath('.//nfe:prod/nfe:CFOP/text()', namespaces=ns)
                    cfop_item = cfop_item[0].strip() if cfop_item else ''
                    
                    adicionar_registro(ncm_item, cfop_item)

            print(" OK")
            arquivos_com_sucesso += 1
        except Exception as e:
            print(f" ERRO ({str(e)})")
            arquivos_com_erro += 1


    print("\n" + "="*50)
    print(f"Varredura concluída!")
    print(f"Sucesso: {arquivos_com_sucesso} arquivos")
    print(f"Falhas: {arquivos_com_erro} arquivos")
    print("="*50 + "\n")

    if not dados_extraidos:
        print("Nenhum dado pôde ser extraído dos XMLs.")
        return

    # Criando o DataFrame do Pandas
    df = pd.DataFrame(dados_extraidos)
    nome_excel = 'planilha_nfe_extraida.xlsx'
    
    print(f"Gerando a planilha Excel '{nome_excel}'...")
    
    # Garantir tratamento de strings em colunas sensíveis (para evitar truncamento de zeros)
    colunas_texto = ['Chave da NFe', 'Número', 'CNPJ Emitente', 'CNPJ Destinatário', 'NCM', 'CFOP', 'Tipo NF', 'Mod Frete']
    for col in colunas_texto:
        if col in df.columns:
            df[col] = df[col].astype(str).replace('nan', '')

    # Salva usando Pandas
    df.to_excel(nome_excel, index=False)

    # Estilizando a planilha usando openpyxl para visual premium
    try:
        wb = load_workbook(nome_excel)
        ws = wb.active
        
        # Ativa linhas de grade na visualização do Excel
        ws.views.sheetView[0].showGridLines = True
        
        # Definição de cores e estilos
        cor_cabecalho = PatternFill(start_color="2A4D69", end_color="2A4D69", fill_type="solid") # Azul Petróleo Premium
        fonte_cabecalho = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        fonte_dados = Font(name="Segoe UI", size=10)
        
        alinhamento_centro = Alignment(horizontal="center", vertical="center")
        alinhamento_esquerda = Alignment(horizontal="left", vertical="center")
        alinhamento_direita = Alignment(horizontal="right", vertical="center")
        
        borda_fina = Border(
            left=Side(style='thin', color='D3D3D3'),
            right=Side(style='thin', color='D3D3D3'),
            top=Side(style='thin', color='D3D3D3'),
            bottom=Side(style='thin', color='D3D3D3')
        )
        
        # Mapeamento do alinhamento desejado por coluna
        alinhamentos = {
            'Chave da NFe': alinhamento_centro,
            'Número': alinhamento_centro,
            'Data Emissão': alinhamento_centro,
            'CNPJ Emitente': alinhamento_centro,
            'CNPJ Destinatário': alinhamento_centro,
            'Natureza Operação': alinhamento_esquerda,
            'Tipo NF': alinhamento_centro,
            'Valor Frete': alinhamento_direita,
            'Mod Frete': alinhamento_centro,
            'Qtd Volume': alinhamento_centro,
            'Espécie': alinhamento_esquerda,
            'NCM': alinhamento_centro,
            'CFOP': alinhamento_centro,
            'Arquivo Origem': alinhamento_esquerda
        }

        # Formatar cabeçalho
        ws.row_dimensions[1].height = 28
        for col_idx, col_name in enumerate(df.columns, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = cor_cabecalho
            cell.font = fonte_cabecalho
            cell.alignment = alinhamento_centro
            cell.border = borda_fina

        # Formatar dados e aplicar estilos adequados
        for r_idx in range(2, ws.max_row + 1):
            ws.row_dimensions[r_idx].height = 20
            for c_idx, col_name in enumerate(df.columns, 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.font = fonte_dados
                cell.border = borda_fina
                
                # Aplica o alinhamento configurado
                if col_name in alinhamentos:
                    cell.alignment = alinhamentos[col_name]
                
                # Força formato texto '@' para campos sensíveis (preserva zeros e evita notação científica)
                if col_name in colunas_texto:
                    cell.number_format = '@'
                # Formata valor decimal para moeda
                elif col_name == 'Valor Frete':
                    cell.number_format = 'R$ #,##0.00'
                # Formata data para exibição limpa (caso esteja no formato ISO NFe)
                elif col_name == 'Data Emissão' and cell.value:
                    val_str = str(cell.value)
                    if 'T' in val_str:
                        # Limpa formato 2026-02-11T15:16:56-03:00 para 2026-02-11 15:16:56
                        val_str = val_str.replace('T', ' ')
                        if '-' in val_str and val_str.count('-') == 3: # possui fuso
                            val_str = val_str.rsplit('-', 1)[0]
                        cell.value = val_str

        # Ajuste automático do tamanho das colunas de acordo com o conteúdo
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            col_name = str(col[0].value)
            
            for cell in col:
                val = str(cell.value or '')
                if cell.row == 1:
                    # Cabeçalhos mais destacados recebem margem maior
                    max_len = max(max_len, len(val) + 4)
                else:
                    max_len = max(max_len, len(val))
            
            # Adiciona folga proporcional ao tipo de conteúdo
            if col_name == 'Chave da NFe':
                ws.column_dimensions[col_letter].width = 48
            elif col_name in ['CNPJ Emitente', 'CNPJ Destinatário']:
                ws.column_dimensions[col_letter].width = 20
            elif col_name == 'Data Emissão':
                ws.column_dimensions[col_letter].width = 22
            else:
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 40)

        wb.save(nome_excel)
        print(f"Formatação visual premium aplicada com sucesso!")
        
    except Exception as e:
        print(f"Aviso: Não foi possível aplicar estilizações avançadas ao arquivo Excel ({str(e)}).")

    end_time = time.time()
    elapsed = end_time - start_time
    print(f"\nProcesso concluído com sucesso!")
    print(f"Total de registros exportados: {len(df)}")
    print(f"Planilha salva em: '{nome_excel}'")
    print(f"Tempo total de execução: {elapsed:.2f} segundos\n")

if __name__ == '__main__':
    extrair_dados_nfe()
